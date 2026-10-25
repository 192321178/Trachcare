#!/usr/bin/env python3
"""Build TrachCare_Poster.pptx - a single-slide academic poster in the same
layout as the BreathTrack reference: title bar, 3-column content grid,
wireframe/prototype rows, footer with QR code + student photo."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

FIGS = os.path.join(os.path.dirname(__file__), 'figs')
OUT = os.path.join(os.path.dirname(__file__), 'TrachCare_Poster.pptx')

TEAL = RGBColor(0x0F, 0x76, 0x6E)
TEAL_DK = RGBColor(0x11, 0x5E, 0x59)
TEAL_LT = RGBColor(0xCC, 0xFB, 0xF1)
MINT = RGBColor(0xEC, 0xFD, 0xF5)
INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x47, 0x55, 0x69)
LINE = RGBColor(0xCB, 0xD5, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(20)
prs.slide_height = Inches(15)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank


def rect(x, y, w, h, fill=WHITE, line=LINE, radius=True, line_w=1.25):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        try:
            shp.adjustments[0] = 0.06
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def text(x, y, w, h, runs, size=13, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, leading=1.15):
    """runs: str, or list of paragraphs; each paragraph str or list of
    (txt, bold, color, size) tuples."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    if isinstance(runs, str):
        runs = [runs]
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = leading
        if isinstance(para, str):
            para = [(para, bold, color, size)]
        for t, b, c, s in para:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(s)
            r.font.bold = b
            r.font.color.rgb = c
            r.font.name = 'Segoe UI'
    return tb


def section(x, y, w, h, title, title_size=15):
    rect(x, y, w, h)
    text(x + 0.15, y + 0.08, w - 0.3, 0.3, title, size=title_size,
         color=TEAL_DK, bold=True)
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                Inches(x + 0.15), Inches(y + 0.42),
                                Inches(w - 0.3), Pt(2))
    ln.fill.solid()
    ln.fill.fore_color.rgb = TEAL
    ln.line.fill.background()
    ln.shadow.inherit = False
    return y + 0.5


def pic(path, x, y, h=None, w=None):
    kw = {}
    if h:
        kw['height'] = Inches(h)
    if w:
        kw['width'] = Inches(w)
    return slide.shapes.add_picture(os.path.join(FIGS, path),
                                    Inches(x), Inches(y), **kw)


# ═══════════════ TITLE BAR ═══════════════
rect(0.35, 0.3, 19.3, 0.95, fill=None, line=None, radius=False)
logo = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.45), Inches(0.42),
                              Inches(0.7), Inches(0.7))
logo.fill.solid(); logo.fill.fore_color.rgb = TEAL
logo.line.color.rgb = TEAL_DK; logo.line.width = Pt(2)
logo.shadow.inherit = False
tf = logo.text_frame
tf.text = 'T'
tf.paragraphs[0].alignment = PP_ALIGN.CENTER
r = tf.paragraphs[0].runs[0]
r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = WHITE
r.font.name = 'Segoe UI'
text(1.35, 0.38, 18, 0.6,
     'TrachCare: A Tracheostomy Care and Monitoring App',
     size=34, color=INK, bold=True)
text(1.35, 0.93, 18, 0.3,
     'Guided tracheostomy care at home — connecting patients, doctors and '
     'admins on one platform', size=14, color=MUTED)
ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(1.28),
                            Inches(19.3), Pt(2.5))
ln.fill.solid(); ln.fill.fore_color.rgb = TEAL
ln.line.fill.background(); ln.shadow.inherit = False

X1, X2, X3, CW = 0.35, 6.85, 13.35, 6.3
TOP = 1.45

# ═══════════════ COLUMN 1 ═══════════════
cy = section(X1, TOP, CW, 1.5, 'Concept')
text(X1 + 0.15, cy + 0.02, CW - 0.3, 0.95, [
    [('TrachCare', True, TEAL_DK, 12.5),
     (' is a tracheostomy self-management app: patients log daily vitals '
      '(RR, HR, SpO2, tube and surgery status) from their phone; doctors '
      'review every report and see flagged abnormal entries; admins manage '
      'accounts and a therapy-video library. Patient, Doctor and Admin '
      'roles share one backend.', False, INK, 12.5)],
], leading=1.12)

cy = section(X1, 3.05, CW, 6.25, 'Competitive Analysis')
apps = [('MyTherapy', 'Medisafe'), ('CareClinic', 'Patient.info')]
ay = cy + 0.1
for row in apps:
    for i, a in enumerate(row):
        rect(X1 + 0.35 + i * 2.85, ay, 2.65, 0.5, fill=MINT, line=LINE)
        text(X1 + 0.35 + i * 2.85, ay + 0.12, 2.65, 0.3, a, size=13,
             color=TEAL_DK, bold=True, align=PP_ALIGN.CENTER)
    ay += 0.62
text(X1 + 0.15, ay + 0.05, CW - 0.3, 4.4, [
    [('1. MyTherapy — ', True, INK, 12),
     ('medication reminders and a general diary; no procedure-specific '
      'tracking or doctor review loop.', False, INK, 12)],
    [('2. Medisafe — ', True, INK, 12),
     ('strong pill reminders, but medication-focused with no clinician '
      'monitoring.', False, INK, 12)],
    [('3. CareClinic — ', True, INK, 12),
     ('broad symptom and measurement diary; nothing specific to '
      'tracheostomy care.', False, INK, 12)],
    [('4. Patient.info — ', True, INK, 12),
     ('education content only; no reporting or connection to the care '
      'team.', False, INK, 12)],
    '',
    [("TrachCare's edge: ", True, TEAL_DK, 12.5),
     ('the only app pairing a daily tracheostomy-specific vitals report '
      'with a doctor review workflow and a guided therapy-video library — '
      'purpose-built for post-surgical home care.', False, INK, 12.5)],
], leading=1.18)

# ═══════════════ COLUMN 2 ═══════════════
cy = section(X2, TOP, CW, 2.35, 'Methods')
text(X2 + 0.15, cy + 0.02, CW - 0.3, 1.8, [
    [('\u2022 ', False, TEAL, 12), ('Daily health report: ', True, INK, 12),
     ('vitals, tube status, comorbidities and labs.', False, INK, 12)],
    [('\u2022 ', False, TEAL, 12), ('Doctor monitoring: ', True, INK, 12),
     ('patient list, report review, flagged alerts.', False, INK, 12)],
    [('\u2022 ', False, TEAL, 12), ('Secure onboarding: ', True, INK, 12),
     ('role login, OTP email verification, admin approval.', False, INK, 12)],
    [('\u2022 ', False, TEAL, 12), ('Education library: ', True, INK, 12),
     ('admin-published therapy videos streamed in-app.', False, INK, 12)],
    [('\u2022 ', False, TEAL, 12), ('Spigotting tracker: ', True, INK, 12),
     ('weaning status visible to patient and doctor.', False, INK, 12)],
], leading=1.12)

cy = section(X2, 3.95, CW, 1.5, 'Problem Statement')
text(X2 + 0.15, cy + 0.02, CW - 0.3, 0.95,
     'Discharged tracheostomy patients must perform daily care and watch for '
     'warning signs — today tracked on paper. Entries are missed, doctors '
     'see patients only at reviews weeks apart, and deterioration is '
     'invisible until it becomes an emergency.', size=12, leading=1.12)

cy = section(X2, 5.55, CW, 1.6, 'User Persona')
face = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(X2 + 0.2),
                              Inches(cy + 0.1), Inches(0.95), Inches(0.95))
face.fill.solid(); face.fill.fore_color.rgb = RGBColor(0x94, 0xA3, 0xB8)
face.line.fill.background(); face.shadow.inherit = False
tf = face.text_frame; tf.text = 'S'
tf.paragraphs[0].alignment = PP_ALIGN.CENTER
r = tf.paragraphs[0].runs[0]
r.font.size = Pt(32); r.font.bold = True; r.font.color.rgb = WHITE
text(X2 + 1.3, cy + 0.08, CW - 1.5, 1.4, [
    [('Name: ', True, INK, 12), ('Santhosh', False, INK, 12),
     ('   Age: ', True, INK, 12), ('58', False, INK, 12)],
    [('Post-surgical tracheostomy patient recovering at home.',
      False, INK, 12)],
    [('Pain points: ', True, INK, 12),
     ('unsure what to record daily, forgets care steps, cannot reach the '
      'doctor quickly, no feedback on readings.', False, INK, 12)],
], leading=1.12)

cy = section(X2, 7.3, CW, 1.15, 'Solution')
text(X2 + 0.15, cy + 0.02, CW - 0.3, 0.6,
     'Replaces the paper diary with a guided daily check-in: structured '
     'vitals go straight to the treating doctor, abnormal reports are '
     'flagged, and training videos keep care correct — closing the gap '
     'between hospital and home.', size=12, leading=1.12)

cy = section(X2, 8.55, CW, 0.75, 'Business Potential')
text(X2 + 0.15, cy + 0.02, CW - 0.3, 0.25,
     'Hospitals & ENT departments · home-care providers · clinic '
     'subscriptions · extends to other post-surgical pathways.',
     size=11.5, leading=1.1)

# ═══════════════ COLUMN 3 ═══════════════
cy = section(X3, TOP, CW, 2.1, 'User Flow Diagram')
pw = pic('fig_flow.png', X3 + 0.55, cy + 0.05, h=1.45)

cy = section(X3, 3.75, CW, 2.55, 'Programming Specifications')
specs = [('Frontend', 'Flutter / Dart (Android)'),
         ('Backend', 'PHP 8.2 REST API (JSON)'),
         ('Database', 'MariaDB / MySQL "trachcare"'),
         ('Server', 'Apache — XAMPP'),
         ('Auth', 'Role login + OTP + approval'),
         ('Media', 'media_kit video · uploads \u2264250 MB'),
         ('Release', 'Signed APK/AAB · targetSdk 36')]
sy = cy + 0.05
for k, v in specs:
    text(X3 + 0.2, sy, 1.9, 0.25, k, size=11.5, color=TEAL_DK, bold=True)
    text(X3 + 2.15, sy, 4.0, 0.25, v, size=11.5, color=INK)
    sy += 0.27

cy = section(X3, 6.45, CW, 2.85, 'Project Timeline')
pic('fig_timeline.png', X3 + 0.3, cy + 0.3, w=5.7)

# ═══════════════ ROW 2: WIREFRAMES / LO-FI ═══════════════
cy = section(0.35, 9.45, 9.4, 2.15, 'Paper Wireframes')
pic('fig_5_2_wireframes.png', 1.5, cy + 0.03, h=1.55)

cy = section(9.9, 9.45, 9.75, 2.15, 'Low Fidelity Prototype')
for i, f in enumerate(['fig_7_2_login.png', 'fig_7_4_daily_update.png',
                       'fig_7_6_add_patient.png']):
    pic(f, 11.6 + i * 1.35, cy + 0.02, h=1.55)
text(11.6, cy + 1.62, 4.4, 0.25,
     'Login  ·  Daily report  ·  Patient chart', size=9, color=MUTED,
     align=PP_ALIGN.CENTER)

# ═══════════════ ROW 3: HI-FI / FINAL ═══════════════
cy = section(0.35, 11.75, 19.3, 2.1,
             'High Fidelity Prototyping  \u2192  Final Product')
hifi = ['fig_7_1_splash.png', 'fig_7_3_patient_dash.png',
        'fig_7_5_doctor_dash.png', 'fig_7_7_approvals.png',
        'fig_7_8_videos.png']
start = 5.4
for i, f in enumerate(hifi):
    pic(f, start + i * 1.35, cy + 0.02, h=1.5)
text(5.4, cy + 1.58, 6.7, 0.25,
     'Splash  ·  Patient dashboard  ·  Doctor monitor  ·  Approvals  ·  '
     'Video library', size=9, color=MUTED, align=PP_ALIGN.CENTER)

# ═══════════════ FOOTER ═══════════════
rect(0.35, 13.95, 19.3, 0.85, fill=TEAL_DK, line=None)
text(0.55, 14.18, 1.3, 0.55, 'Scan For\nDetails', size=11, color=WHITE,
     bold=True, leading=1.0)
# white backing for QR
rect(1.9, 13.99, 0.77, 0.77, fill=WHITE, line=None)
pic('qr_code.png', 1.93, 14.02, h=0.71)
text(3.0, 14.18, 6.5, 0.45,
     'TrachCare — tracheostomy care, guided at home\nPatient · Doctor · '
     'Admin  |  Flutter + PHP + MySQL', size=11, color=WHITE, leading=1.1)

# student photo + details (right side)
pic('student_photo.png', 18.6, 13.62, h=1.18)
text(13.4, 14.05, 5.0, 0.7, [
    [('Sivanesan', True, WHITE, 15)],
    [('IT Department · SIMATS', False, WHITE, 11.5)],
    [('Reg. No. __________ · email@saveetha.com', False, WHITE, 11.5)],
], align=PP_ALIGN.RIGHT, leading=1.1)

prs.save(OUT)
print('written:', OUT)
